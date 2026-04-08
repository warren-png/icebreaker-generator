import streamlit as st
import anthropic
import base64
import os
import io
from pathlib import Path
from dotenv import load_dotenv
from utils.auth import check_password

load_dotenv()

st.set_page_config(
    page_title="Scorecard | Biz Dev Entourage",
    page_icon="🎯",
    layout="wide"
)

# — Authentification —
if not check_password():
    st.stop()

# ---------------------------------------------------------------------------
# COMMERCIAUX
# ---------------------------------------------------------------------------

COMMERCIAUX = {
    "Warren": {
        "linkedin": "https://www.linkedin.com/in/warren-elbaz/",
        "tel": "06 50 60 22 61"
    },
    "Helder": {
        "linkedin": "https://www.linkedin.com/in/helder-alturas-48010463/",
        "tel": "06 22 30 96 11"
    }
}

# ---------------------------------------------------------------------------
# LOGO — resize via Pillow (cached), sidebar upload en override
# ---------------------------------------------------------------------------

@st.cache_resource
def load_logo_base64_cached() -> str | None:
    """Charge le logo depuis le disque, le redimensionne et le met en cache."""
    candidates = [
        Path(__file__).parent.parent / "logo_entourage.png",
        Path(__file__).parent / "logo_entourage.png",
        Path(os.getcwd()) / "logo_entourage.png",
    ]
    for path in candidates:
        if path.exists():
            try:
                from PIL import Image
                img = Image.open(path)
                img.thumbnail((600, 120), Image.LANCZOS)
                buf = io.BytesIO()
                img.save(buf, format="PNG", optimize=True)
                return base64.b64encode(buf.getvalue()).decode()
            except ImportError:
                with open(path, "rb") as f:
                    return base64.b64encode(f.read()).decode()
    return None


def get_logo_base64() -> str | None:
    """Retourne le logo en base64 (sidebar upload prioritaire, sinon disque)."""
    if st.session_state.get("logo_b64"):
        return st.session_state.logo_b64
    cached = load_logo_base64_cached()
    if cached:
        st.session_state.logo_b64 = cached
    return cached


# ---------------------------------------------------------------------------
# EXTRACTION TEXTE
# ---------------------------------------------------------------------------

def extract_text(uploaded_file) -> str | None:
    name = uploaded_file.name.lower()

    if name.endswith(".txt"):
        return uploaded_file.read().decode("utf-8", errors="ignore")

    elif name.endswith(".pdf"):
        try:
            import pdfplumber
            with pdfplumber.open(io.BytesIO(uploaded_file.read())) as pdf:
                pages = [page.extract_text() or "" for page in pdf.pages]
            return "\n".join(pages)
        except ImportError:
            st.error("pdfplumber non installé : `pip install pdfplumber`")
            return None

    elif name.endswith(".docx"):
        try:
            from docx import Document
            doc = Document(io.BytesIO(uploaded_file.read()))
            return "\n".join(p.text for p in doc.paragraphs)
        except ImportError:
            st.error("python-docx non installé : `pip install python-docx`")
            return None

    st.error("Format non supporté.")
    return None


# ---------------------------------------------------------------------------
# INJECTION METADATA (logo + client + commercial) — APRÈS génération Claude
# ---------------------------------------------------------------------------

def inject_metadata(html: str, client_name: str, commercial: str) -> str:
    """Injecte le logo, le nom client et le footer commercial dans le HTML généré."""
    logo_b64 = get_logo_base64()

    # Logo
    if logo_b64:
        logo_tag = (
            f'<img src="data:image/png;base64,{logo_b64}" '
            f'alt="Entourage Recrutement" style="height:18mm; object-fit:contain;">'
        )
    else:
        logo_tag = '<span style="color:white;font-weight:800;font-size:13pt;letter-spacing:1px;">ENTOURAGE RECRUTEMENT</span>'
    html = html.replace("{{LOGO}}", logo_tag)

    # Nom client
    html = html.replace("{{NOM_CLIENT}}", client_name.upper() if client_name else "")

    # Footer commercial
    info = COMMERCIAUX[commercial]
    footer_html = (
        f'<a href="{info["linkedin"]}" target="_blank">{commercial}</a>'
        f' &nbsp;{info["tel"]}'
    )
    html = html.replace("{{FOOTER_COMMERCIAL}}", footer_html)

    return html


# ---------------------------------------------------------------------------
# TEMPLATE HTML SCORECARD
# ---------------------------------------------------------------------------

SCORECARD_TEMPLATE = """<!DOCTYPE html>
<html lang="fr">
<head>
    <meta charset="UTF-8">
    <title>Brief de Mission</title>
    <link rel="preconnect" href="https://fonts.googleapis.com">
    <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
    <link href="https://fonts.googleapis.com/css2?family=Manrope:wght@400;500;700;800&family=Playfair+Display:ital,wght@0,400;0,700;1,400&display=swap" rel="stylesheet">
    <link rel="stylesheet" href="https://cdnjs.cloudflare.com/ajax/libs/font-awesome/6.5.1/css/all.min.css">
    <style>
        * { box-sizing: border-box; margin: 0; padding: 0; -webkit-print-color-adjust: exact; print-color-adjust: exact; }
        body {
            background-color: #555;
            font-family: 'Manrope', sans-serif;
            display: flex;
            flex-direction: column;
            align-items: center;
            padding: 40px 0;
        }
        .page {
            width: 210mm;
            height: 297mm;
            max-height: 297mm;
            overflow: hidden;
            background: white;
            box-shadow: 0 0 20px rgba(0,0,0,0.5);
            position: relative;
            display: flex;
            flex-direction: column;
        }
        @media print {
            @page { size: A4 portrait; margin: 0; }
            body { background: none; padding: 0; }
            .page { margin: 0; box-shadow: none; width: 210mm; height: 297mm; }
        }
        .header {
            height: 25mm;
            background: #000;
            display: flex;
            justify-content: space-between;
            align-items: center;
            padding: 0 15mm;
            border-bottom: 2mm solid #FFD700;
            flex-shrink: 0;
        }
        .brand { display: flex; align-items: center; }
        .brand img { height: 18mm; object-fit: contain; }
        .doc-title { color: #fff; font-size: 10pt; text-transform: uppercase; letter-spacing: 1px; font-weight: 600; text-align: right; }
        .doc-title .mandat { display: block; font-size: 7pt; color: #FFD700; margin-top: 3px; }
        .doc-title .client { display: block; font-size: 8pt; color: #fff; opacity: 0.85; margin-top: 2px; letter-spacing: 0.5px; }
        .content { padding: 6mm 15mm 5mm; flex-grow: 1; display: flex; flex-direction: column; justify-content: space-between; }
        .section-title {
            font-family: 'Playfair Display', serif;
            font-size: 13pt;
            color: #000;
            border-bottom: 1px solid #eee;
            padding-bottom: 1.5mm;
            margin-bottom: 4mm;
            display: flex;
            align-items: center;
            gap: 8px;
        }
        .section-title i { color: #FFD700; font-size: 11pt; }
        .summary-box {
            background: #f8f9fa;
            border-left: 3mm solid #000;
            padding: 4mm;
            display: grid;
            grid-template-columns: 1fr 1fr 1fr;
            gap: 4mm;
        }
        .sum-col h4 { font-family: 'Playfair Display', serif; font-size: 10pt; margin-bottom: 1.5mm; color: #000; border-bottom: 1px solid #FFD700; display: inline-block; text-transform: uppercase; }
        .sum-col p { font-size: 8.5pt; line-height: 1.4; color: #444; text-align: justify; }
        .score-table { width: 100%; border-collapse: collapse; font-size: 8.5pt; }
        .score-table th { text-align: left; padding: 2mm 3mm; background: #000; color: #fff; text-transform: uppercase; font-size: 8pt; letter-spacing: 0.5px; }
        .score-table td { padding: 2.5mm 3mm; border-bottom: 1px solid #eee; vertical-align: middle; }
        .score-cat { width: 25%; font-weight: 800; color: #000; border-right: 2px solid #FFD700; text-transform: uppercase; }
        .score-weight { width: 15%; text-align: center; font-weight: 800; color: #000; font-size: 9pt; background-color: #fcfcfc; border-right: 1px solid #eee; }
        .score-desc { color: #444; line-height: 1.3; padding-left: 15px !important; }
        .score-desc strong { color: #000; }
        .process-container {
            display: flex;
            justify-content: space-between;
            align-items: flex-start;
            margin-top: 2mm;
            position: relative;
        }
        .process-container::before {
            content: ''; position: absolute; top: 7mm; left: 20px; right: 20px; height: 1px; background: #ddd; z-index: 0;
        }
        .step { position: relative; z-index: 1; text-align: center; width: 22%; }
        .step-num {
            width: 11mm; height: 11mm; background: #000; color: #FFD700;
            border-radius: 50%; display: flex; align-items: center; justify-content: center;
            font-weight: 800; font-size: 9pt; margin: 0 auto 2mm auto;
            border: 2px solid #fff; box-shadow: 0 0 0 1px #000;
        }
        .step-title { font-weight: 800; font-size: 8pt; margin-bottom: 1px; text-transform: uppercase; }
        .step-who { font-size: 7.5pt; color: #555; font-style: italic; }
        .salary-box {
            margin-top: 4mm;
            text-align: center;
            border: 1px dashed #ccc;
            padding: 2.5mm;
            border-radius: 4px;
            background: #fffcf5;
            font-family: 'Playfair Display', serif;
            font-size: 10pt;
            color: #000;
        }
        .footer {
            height: 10mm;
            background: #f8f9fa;
            border-top: 1px solid #ddd;
            display: flex;
            justify-content: center;
            align-items: center;
            font-size: 8pt;
            color: #555;
            margin-top: auto;
            gap: 6px;
        }
        .footer a {
            color: #000;
            font-weight: 800;
            text-decoration: none;
            border-bottom: 2px solid #FFD700;
        }
    </style>
</head>
<body>
    <div class="page">
        <div class="header">
            <div class="brand">
                {{LOGO}}
            </div>
            <div class="doc-title">
                BRIEF DE MISSION
                <span class="mandat">MANDAT : {{TITRE_POSTE}}</span>
                <span class="client">{{NOM_CLIENT}}</span>
            </div>
        </div>

        <div class="content">
            <div>
                <h2 class="section-title"><i class="fa-solid fa-bullseye"></i> Vision & Contexte</h2>
                <div class="summary-box">
                    <div class="sum-col">
                        <h4>{{TITRE_BLOC_1}}</h4>
                        <p>{{TEXTE_BLOC_1}}</p>
                    </div>
                    <div class="sum-col">
                        <h4>{{TITRE_BLOC_2}}</h4>
                        <p>{{TEXTE_BLOC_2}}</p>
                    </div>
                    <div class="sum-col">
                        <h4>{{TITRE_BLOC_3}}</h4>
                        <p>{{TEXTE_BLOC_3}}</p>
                    </div>
                </div>
            </div>

            <div>
                <h2 class="section-title"><i class="fa-solid fa-star"></i> Scorecard</h2>
                <table class="score-table">
                    <thead>
                        <tr>
                            <th>Critère Clé</th>
                            <th style="text-align:center">Pondération</th>
                            <th>Définition du Succès (Indicateurs)</th>
                        </tr>
                    </thead>
                    <tbody>
                        {{LIGNES_SCORECARD}}
                    </tbody>
                </table>
            </div>

            <div>
                <h2 class="section-title"><i class="fa-solid fa-users-viewfinder"></i> Processus de Recrutement</h2>
                <div class="process-container">
                    {{ETAPES_PROCESSUS}}
                </div>
                <div class="salary-box">
                    <strong>PACKAGE CIBLE :</strong> {{PACKAGE}}
                </div>
            </div>
        </div>

        <div class="footer">
            {{FOOTER_COMMERCIAL}}
        </div>
    </div>
</body>
</html>"""


# ---------------------------------------------------------------------------
# GENERATION CLAUDE (sans logo — injecté après)
# ---------------------------------------------------------------------------

SYSTEM_PROMPT = """Tu es expert en recrutement de cadres dirigeants pour Entourage Recrutement, cabinet de chasse de têtes spécialisé en finance (DAF, CFO, M&A, contrôle de gestion).

Tu génères des Briefs de Mission (scorecards) professionnels à partir de retranscriptions d'appels de qualification client.

Règles absolues :
- Retourne UNIQUEMENT le code HTML complet, sans explication, sans balises markdown
- Respecte scrupuleusement la structure du template fourni
- Ne modifie JAMAIS les placeholders {{LOGO}}, {{NOM_CLIENT}}, {{FOOTER_COMMERCIAL}} — laisse-les exactement tels quels
- Les pondérations de la scorecard doivent totaliser exactement 100%
- Entre 4 et 5 critères dans la scorecard (pas plus)
- Utilise un langage professionnel, précis et orienté résultats
- Si une info n'est pas explicite dans la retranscription, déduis-la intelligemment du contexte

CONTRAINTE DE FORMAT STRICTE — PRIORITÉ ABSOLUE :
- Le document doit tenir sur UNE SEULE page A4 (210mm × 297mm). Tout débordement sera coupé.
- Blocs "Vision & Contexte" (sum-col) : 3 phrases maximum par bloc, 280 caractères maximum chacun.
- Scorecard : 4 à 5 critères maximum. Chaque description de succès : 1 à 2 lignes, 180 caractères maximum.
- Processus de recrutement : exactement 4 étapes. Titres de 2-4 mots. Intervenants en 2-3 mots.
- Package : une seule ligne synthétique (ex: "90-110K€ fixe + 15-20% variable + BSPCE").
- Langage professionnel et précis, orienté résultats. Pas de répétitions entre les sections."""


def generate_scorecard(
    transcription: str,
    modification: str = None,
    previous_html: str = None
) -> str:
    client = anthropic.Anthropic(api_key=os.getenv("ANTHROPIC_API_KEY"))

    if modification and previous_html:
        user_content = f"""Voici la scorecard HTML que tu as générée :

{previous_html}

L'utilisateur demande la modification suivante :
{modification}

Applique cette modification et retourne le HTML complet mis à jour.
IMPORTANT : Ne modifie pas les placeholders {{{{LOGO}}}}, {{{{NOM_CLIENT}}}}, {{{{FOOTER_COMMERCIAL}}}} s'ils sont présents.
Retourne UNIQUEMENT le HTML, sans explication."""
    else:
        user_content = f"""Voici la retranscription d'un appel de qualification client :

---
{transcription}
---

Remplis ce template HTML de scorecard avec les informations extraites de la retranscription.
IMPORTANT : Laisse les placeholders {{{{LOGO}}}}, {{{{NOM_CLIENT}}}}, {{{{FOOTER_COMMERCIAL}}}} exactement tels quels dans le HTML.

TEMPLATE À REMPLIR :
{SCORECARD_TEMPLATE}

Retourne le HTML complet avec tous les autres placeholders {{{{...}}}} remplacés par les vraies informations."""

    message = client.messages.create(
        model="claude-sonnet-4-5",
        max_tokens=8096,
        system=SYSTEM_PROMPT,
        messages=[{"role": "user", "content": user_content}]
    )

    import re as _re
    html = message.content[0].text.strip()
    html = _re.sub(r"^```[^\n]*\n", "", html)
    html = _re.sub(r"\n```\s*$", "", html.strip())
    return html.strip()


# ---------------------------------------------------------------------------
# EXPORT PDF — via impression navigateur (rendu parfait, zéro dépendance)
# ---------------------------------------------------------------------------

def get_print_button_html(html_content: str, label: str = "📄 Télécharger PDF") -> str:
    """Bouton qui ouvre le document dans une nouvelle fenêtre et déclenche l'impression."""
    b64 = base64.b64encode(html_content.encode("utf-8")).decode()
    return f"""
    <script>
    function printDoc() {{
        var w = window.open('', '_blank');
        var bytes = Uint8Array.from(atob('{b64}'), function(c) {{ return c.charCodeAt(0); }});
        var html = new TextDecoder('utf-8').decode(bytes);
        w.document.open();
        w.document.write(html);
        w.document.close();
        w.onload = function() {{
            setTimeout(function() {{ w.print(); }}, 400);
        }};
    }}
    </script>
    <button onclick="printDoc()" style="
        background-color: #000;
        color: #FFD700;
        border: none;
        padding: 10px 18px;
        font-size: 14px;
        font-weight: 700;
        border-radius: 6px;
        cursor: pointer;
        width: 100%;
        font-family: sans-serif;
        letter-spacing: 0.5px;
    ">{label}</button>
    """


# ---------------------------------------------------------------------------
# UI — SIDEBAR
# ---------------------------------------------------------------------------

with st.sidebar:
    st.markdown("### 🖼 Logo Entourage")
    logo_file = st.file_uploader(
        "Charger le logo (si non détecté)",
        type=["png", "jpg", "jpeg"],
        label_visibility="collapsed",
        key="logo_uploader"
    )
    if logo_file:
        try:
            from PIL import Image
            img = Image.open(logo_file)
            img.thumbnail((600, 120), Image.LANCZOS)
            buf = io.BytesIO()
            img.save(buf, format="PNG", optimize=True)
            st.session_state.logo_b64 = base64.b64encode(buf.getvalue()).decode()
        except ImportError:
            st.session_state.logo_b64 = base64.b64encode(logo_file.read()).decode()
        st.success("Logo chargé ✓")

    logo = get_logo_base64()
    if logo:
        logo_bytes = base64.b64decode(logo)
        st.image(io.BytesIO(logo_bytes), use_container_width=True)
    else:
        st.caption("⚠️ Logo non trouvé. Uploadez-le ci-dessus.")

# ---------------------------------------------------------------------------
# UI — PAGE PRINCIPALE
# ---------------------------------------------------------------------------

st.title("🎯 Générateur de Scorecard")
st.caption("Remplis les informations, uploade la retranscription → Brief de Mission généré automatiquement")

st.divider()

# Informations client et commercial
col1, col2 = st.columns([3, 1])
with col1:
    client_name = st.text_input(
        "Nom du client *",
        placeholder="Ex : TD Williamson, BNP Paribas...",
        help="Apparaît dans l'en-tête du document"
    )
with col2:
    commercial = st.radio("Commercial", list(COMMERCIAUX.keys()), horizontal=False)

st.divider()

# Upload retranscription
uploaded_file = st.file_uploader(
    "Retranscription de l'appel de qualification",
    type=["pdf", "txt", "docx"],
    help="Formats acceptés : PDF, TXT, DOCX"
)

if uploaded_file:
    transcription_text = extract_text(uploaded_file)

    if transcription_text:
        char_count = len(transcription_text)
        st.success(f"✅ Fichier lu — {char_count:,} caractères extraits")

        btn_disabled = not client_name.strip()
        if btn_disabled:
            st.warning("⚠️ Renseigne le nom du client avant de générer.")

        if st.button("🚀 Générer la Scorecard", type="primary", disabled=btn_disabled):
            with st.spinner("Claude analyse la retranscription et génère le Brief de Mission..."):
                try:
                    raw_html = generate_scorecard(transcription_text)
                    final_html = inject_metadata(raw_html, client_name, commercial)
                    st.session_state.scorecard_html = final_html
                    st.session_state.scorecard_raw_html = raw_html
                    st.session_state.scorecard_transcription = transcription_text
                    st.session_state.scorecard_client = client_name
                    st.session_state.scorecard_commercial = commercial
                    st.rerun()
                except Exception as e:
                    st.error(f"Erreur lors de la génération : {e}")

# ---------------------------------------------------------------------------
# RÉSULTAT
# ---------------------------------------------------------------------------

if "scorecard_html" in st.session_state:
    st.divider()

    # Boutons d'action
    col1, col2, col3 = st.columns([2, 2, 6])
    with col1:
        st.components.v1.html(
            get_print_button_html(st.session_state.scorecard_html),
            height=50
        )
    with col2:
        if st.button("🗑️ Réinitialiser", use_container_width=True):
            for k in ["scorecard_html", "scorecard_transcription", "scorecard_client", "scorecard_commercial"]:
                st.session_state.pop(k, None)
            st.rerun()

    st.caption("💡 Cliquez sur 📄 Télécharger PDF → une fenêtre s'ouvre → Fichier → Imprimer → Enregistrer en PDF.")

    # Aperçu
    st.subheader("Aperçu")
    st.components.v1.html(st.session_state.scorecard_html, height=1250, scrolling=True)

    # Zone modifications
    st.divider()
    st.subheader("✏️ Demander des modifications")
    st.caption("Une information manquante ou incorrecte ? Décris la correction ici.")

    modification = st.text_area(
        "Modifications souhaitées",
        placeholder=(
            "Exemples :\n"
            "• \"Ajoute un critère sur l'expérience internationale à 15%\"\n"
            "• \"Le package est 90-100k fixe + 20% variable\"\n"
            "• \"Le processus a 3 étapes : Entourage, DRH, CEO\"\n"
            "• \"Modifie le contexte : l'entreprise est en phase de croissance externe\""
        ),
        height=130,
        label_visibility="collapsed"
    )

    if st.button("🔄 Appliquer les modifications", type="secondary"):
        if modification.strip():
            with st.spinner("Application des modifications..."):
                try:
                    # Sauvegarde de sécurité avant toute modification
                    backup_raw = st.session_state.scorecard_raw_html
                    backup_html = st.session_state.scorecard_html

                    raw_html = generate_scorecard(
                        st.session_state.scorecard_transcription,
                        modification=modification,
                        previous_html=backup_raw
                    )

                    # Validation : HTML non vide et structure minimale présente
                    if not raw_html or len(raw_html) < 200 or "</html>" not in raw_html.lower():
                        st.error("❌ La réponse de Claude est incomplète. Le document original est conservé.")
                        st.stop()

                    final_html = inject_metadata(
                        raw_html,
                        st.session_state.scorecard_client,
                        st.session_state.scorecard_commercial
                    )
                    st.session_state.scorecard_html = final_html
                    st.session_state.scorecard_raw_html = raw_html
                    st.rerun()
                except Exception as e:
                    st.error(f"Erreur : {e}")
        else:
            st.warning("Décris d'abord les modifications souhaitées.")
