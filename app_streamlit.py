"""
Page d'accueil — Mandats : base partagée des mandats en cours et passés.
(Remplace l'ancienne page Icebreaker, retirée en octobre 2026 ; voir l'historique git.)

Pour chaque mandat : nom, entreprise, date, responsable, statut, notes,
fiche de poste (PDF/Word ou texte collé), échanges audio (manager, RH…)
et leur transcription intégrale (AssemblyAI, avec locuteurs et horodatage).

Stockage : Drive partagé Entourage, dossier « Mandats » (voir mandats_store.py).
"""

import io
import os
import re
import sys
import mimetypes
from datetime import date, datetime

import streamlit as st

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from utils.auth import check_password
from utils.ui import inject_global_styles
from mandats_store import StorageError
from utils.mandats_data import store, load_snapshot, fetch_bytes
from utils import leonar


st.set_page_config(page_title="Mandats | Entourage", page_icon="🗂️", layout="wide")
inject_global_styles()

if not check_password():
    st.stop()


RESPONSABLES = ["Warren", "Helder"]          # commerciaux (pied de page scorecard / dossier)
PROPRIETAIRES = ["Warren", "Helder", "Bruno"]  # propriétaires possibles du projet Leonar
STATUTS = ["En cours", "Pourvu", "Clos"]
STATUT_COLORS = {"En cours": "orange", "Pourvu": "green", "Clos": "gray"}
AUDIO_TYPES = ["m4a", "mp3", "wav", "mp4", "aac", "ogg", "webm", "flac", "mpeg", "mov"]
FICHE_TYPES = ["pdf", "docx", "doc"]

st.markdown(
    """
    <style>
      .mandat-meta { color: #71717A; font-size: 0.86rem; margin-top: -0.2rem; }
      .mandat-title { font-weight: 800; font-size: 1.02rem; color: #0A0A0A; letter-spacing: -0.01em; }
      .mandat-files { color: #52525B; font-size: 0.82rem; }
      .mandat-empty {
          text-align: center; padding: 48px 20px; color: #71717A;
          border: 1.5px dashed #D4D4CE; border-radius: 14px; background: #FFFFFF;
      }
      .mandat-empty b { color: #0A0A0A; font-size: 1.05rem; display: block; margin-bottom: 4px; }
      .mandat-transcript { font-size: 0.9rem; line-height: 1.6; }
      .mandat-transcript .ts { color: #A1A1AA; font-variant-numeric: tabular-nums; font-size: 0.8rem; }
      .mandat-transcript .spk { font-weight: 750; color: #0A0A0A; }
      .mandat-transcript p { margin: 0 0 0.55rem 0; }
      /* Sur cette page, beaucoup de téléchargements : style secondaire, plus léger */
      .stDownloadButton > button {
          background: #FFFFFF !important;
          color: #18181B !important;
          border: 1px solid #E4E4E0 !important;
          box-shadow: 0 1px 2px rgba(16,16,16,0.04) !important;
      }
      .stDownloadButton > button:hover {
          border-color: #D4D4CE !important;
          color: #0A0A0A !important;
          background: #FCFCFA !important;
          box-shadow: 0 1px 2px rgba(16,16,16,0.04), 0 6px 20px rgba(16,16,16,0.05) !important;
      }
    </style>
    """,
    unsafe_allow_html=True,
)


# ─────────────────────────────────────────────────────────────────────────────
# Stockage
# ─────────────────────────────────────────────────────────────────────────────


def refresh_and_rerun(flash: str | None = None, mandat_id: str | None = None):
    load_snapshot.clear()
    if flash:
        st.session_state["_mandats_flash"] = flash
    if mandat_id is not None:
        st.query_params["mandat"] = mandat_id
    st.rerun()


def ensure_env(key: str) -> str | None:
    """Les secrets Streamlit Cloud ne sont pas toujours exportés en variables d'env."""
    if not os.getenv(key):
        try:
            if key in st.secrets:
                os.environ[key] = str(st.secrets[key])
        except Exception:
            pass
    return os.getenv(key)


# ─────────────────────────────────────────────────────────────────────────────
# Helpers d'affichage
# ─────────────────────────────────────────────────────────────────────────────


def md(text: str) -> str:
    """Échappe le texte saisi pour l'afficher en markdown."""
    return re.sub(r"([\\`*_\[\]{}<>#|$~])", r"\\\1", str(text or ""))


def fmt_date(iso: str) -> str:
    try:
        return datetime.fromisoformat(iso).strftime("%d/%m/%Y")
    except Exception:
        return iso or "—"


def fmt_datetime(iso: str) -> str:
    try:
        return datetime.fromisoformat(iso.replace("Z", "+00:00")).strftime("%d/%m/%Y à %H:%M")
    except Exception:
        return iso or "—"


def fmt_size(n: int) -> str:
    if not n:
        return ""
    if n < 1024:
        return f"{n} o"
    if n < 1024 * 1024:
        return f"{n / 1024:.0f} Ko"
    return f"{n / (1024 * 1024):.1f} Mo"


def ms_ts(ms: int) -> str:
    s = int((ms or 0) / 1000)
    return f"{s // 3600:02d}:{(s % 3600) // 60:02d}:{s % 60:02d}"


def strip_prefix(name: str) -> str:
    return re.sub(r"^(Audio|Transcription|Fiche de poste) — ", "", name)


def guess_mime(name: str, fallback: str = "application/octet-stream") -> str:
    return mimetypes.guess_type(name)[0] or fallback


def audio_display_name(audio: dict) -> str:
    return audio["props"].get("label") or strip_prefix(audio["name"])


def badges(m: dict) -> str:
    statut = m.get("statut", "En cours")
    color = STATUT_COLORS.get(statut, "gray")
    parts = [f":{color}-badge[{md(statut)}]"]
    if m.get("responsable"):
        parts.append(f":gray-badge[:material/person: {md(m['responsable'])}]")
    parts.append(f":gray-badge[:material/calendar_today: {fmt_date(m.get('date', ''))}]")
    if m.get("leonar_project_id"):
        parts.append(":violet-badge[:material/link: Leonar]")
    return " ".join(parts)


def files_summary(files: list[dict]) -> str:
    n_fiche = sum(1 for f in files if f["kind"] == "fiche")
    n_audio = sum(1 for f in files if f["kind"] == "audio")
    n_tr = sum(1 for f in files if f["kind"] == "transcript")
    bits = [
        "Fiche de poste ✓" if n_fiche else "Pas de fiche",
        f"{n_audio} audio{'s' if n_audio > 1 else ''}",
        f"{n_tr} transcription{'s' if n_tr > 1 else ''}",
    ]
    return " · ".join(bits)


def sort_key(m: dict):
    return (m.get("date", ""), m.get("created_at", ""))


# ─────────────────────────────────────────────────────────────────────────────
# Transcriptions
# ─────────────────────────────────────────────────────────────────────────────

_LINE_RE = re.compile(r"^\[(\d{2}:\d{2}:\d{2})\] (Locuteur [^:]+) : (.*)$")


def build_transcript(mandat: dict, audio: dict, result: dict) -> str:
    duration = ms_ts((result.get("duration_seconds") or 0) * 1000)
    header = [
        "TRANSCRIPTION INTÉGRALE",
        f"Mandat : {mandat.get('nom', '')}",
        f"Entreprise : {mandat.get('entreprise', '')}",
        f"Date du mandat : {fmt_date(mandat.get('date', ''))}",
        f"Audio : {audio_display_name(audio)}",
        f"Durée : {duration} · {result.get('speakers_count', 0)} locuteur(s)",
        "─" * 60,
        "",
    ]
    lines = []
    for u in result.get("utterances") or []:
        text = (u.get("text") or "").strip()
        if text:
            lines.append(f"[{ms_ts(u.get('start_ms'))}] Locuteur {u.get('speaker', '?')} : {text}")
    if not lines:
        lines = [result.get("formatted_text", "")]
    return "\n".join(header + lines) + "\n"


def transcript_to_docx(text: str, title: str) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    doc.add_heading(title, level=1)
    for line in text.splitlines():
        m = _LINE_RE.match(line)
        p = doc.add_paragraph()
        if m:
            ts = p.add_run(f"[{m.group(1)}] ")
            ts.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
            ts.font.size = Pt(9)
            p.add_run(f"{m.group(2)} : ").bold = True
            p.add_run(m.group(3))
        else:
            p.add_run(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_transcript(text: str):
    html = []
    for line in text.splitlines():
        m = _LINE_RE.match(line)
        if m:
            esc = lambda s: s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            html.append(
                f'<p><span class="ts">{m.group(1)}</span>&nbsp; '
                f'<span class="spk">{esc(m.group(2))}</span> — {esc(m.group(3))}</p>'
            )
    if not html:
        st.text(text)
        return
    st.markdown(f'<div class="mandat-transcript">{"".join(html)}</div>', unsafe_allow_html=True)


def transcribe_and_store(mandat: dict, audio: dict, audio_bytes: bytes) -> bool:
    """Transcrit un audio déjà stocké et enregistre la transcription à côté."""
    if not ensure_env("ASSEMBLYAI_API_KEY"):
        st.error("ASSEMBLYAI_API_KEY manquante : l'audio est enregistré, mais pas transcrit.")
        return False
    from coach_prospection.transcription import transcribe_audio_file

    name = audio_display_name(audio)
    with st.status(f"Transcription de « {name} »…", expanded=True) as status:
        st.caption("Gardez cette page ouverte : comptez 2 à 5 minutes pour 1 h d'audio.")
        try:
            result = transcribe_audio_file(
                audio_bytes,
                filename=strip_prefix(audio["name"]),
                language="fr",
                progress_callback=status.write,
            )
            text = build_transcript(mandat, audio, result)
            stem = os.path.splitext(strip_prefix(audio["name"]))[0]
            store().upload(
                mandat["id"],
                f"Transcription — {stem}.txt",
                text.encode("utf-8"),
                "text/plain",
                {"ent_kind": "transcript", "audio_id": audio["id"], "label": audio["props"].get("label")},
            )
        except Exception as e:
            status.update(label=f"Transcription échouée — {name}", state="error")
            st.error(f"La transcription a échoué : {e}. L'audio est bien enregistré, vous pouvez relancer.")
            return False
        status.update(label=f"Transcription enregistrée — {name}", state="complete", expanded=False)
    return True


def add_audio(mandat: dict, uploaded, label: str = "") -> bool:
    data = uploaded.getvalue()
    with st.spinner(f"Envoi de « {uploaded.name} » sur le Drive…"):
        audio = store().upload(
            mandat["id"],
            f"Audio — {uploaded.name}",
            data,
            uploaded.type or guess_mime(uploaded.name, "audio/mpeg"),
            {"ent_kind": "audio", "label": label.strip() or None},
        )
    return transcribe_and_store(mandat, audio, data)


def save_fiche(mandat_id: str, files: list[dict], uploaded=None, text: str = ""):
    fiches = [f for f in files if f["kind"] == "fiche"]
    if uploaded is not None:
        name = f"Fiche de poste — {uploaded.name}"
        mime = uploaded.type or guess_mime(uploaded.name)
        existing = next((f for f in fiches if f["props"].get("format") == "file"), None)
        if existing:
            store().replace_content(existing["id"], name, uploaded.getvalue(), mime)
        else:
            store().upload(mandat_id, name, uploaded.getvalue(), mime, {"ent_kind": "fiche", "format": "file"})
    if text.strip():
        name = "Fiche de poste (texte).txt"
        existing = next((f for f in fiches if f["props"].get("format") == "text"), None)
        if existing:
            store().replace_content(existing["id"], name, text.strip().encode("utf-8"), "text/plain")
        else:
            store().upload(
                mandat_id, name, text.strip().encode("utf-8"), "text/plain",
                {"ent_kind": "fiche", "format": "text"},
            )


# ─────────────────────────────────────────────────────────────────────────────
# Formulaires
# ─────────────────────────────────────────────────────────────────────────────


def init_mandat_fields(prefix: str, m: dict | None = None):
    """Valeurs initiales des champs (avant leur affichage) ; n'écrase pas une saisie en cours."""
    m = m or {}
    try:
        d = date.fromisoformat(m.get("date", "")) if m.get("date") else date.today()
    except ValueError:
        d = date.today()
    resp = m.get("responsable") if m.get("responsable") in RESPONSABLES else RESPONSABLES[0]
    proprios = [p for p in (m.get("proprietaires") or [resp]) if p in PROPRIETAIRES]
    defaults = {
        "nom": m.get("nom", ""),
        "ent": m.get("entreprise", ""),
        "date": d,
        "resp": resp,
        "statut": m.get("statut") if m.get("statut") in STATUTS else STATUTS[0],
        "proprio": proprios,
        "notes": m.get("notes", ""),
    }
    for k, v in defaults.items():
        st.session_state.setdefault(f"{prefix}_{k}", v)


def mandat_fields(prefix: str, m: dict | None = None) -> dict:
    init_mandat_fields(prefix, m)
    c1, c2 = st.columns(2)
    nom = c1.text_input("Nom du mandat *", placeholder="Ex : Contrôleur de gestion H/F", key=f"{prefix}_nom")
    entreprise = c2.text_input("Entreprise *", placeholder="Ex : TD Williamson", key=f"{prefix}_ent")
    c3, c4, c5 = st.columns(3)
    the_date = c3.date_input("Date", format="DD/MM/YYYY", key=f"{prefix}_date")
    responsable = c4.selectbox("Responsable", RESPONSABLES, key=f"{prefix}_resp")
    statut = c5.selectbox("Statut", STATUTS, key=f"{prefix}_statut")
    proprietaires = st.multiselect(
        "Propriétaires du projet Leonar", PROPRIETAIRES, key=f"{prefix}_proprio",
        help="Membres propriétaires du projet dans Leonar (le responsable commercial, et le sourceur si besoin).",
    )
    notes = st.text_area(
        "Notes internes", height=100,
        placeholder="Contexte, points d'attention, interlocuteurs…", key=f"{prefix}_notes",
    )
    return {
        "nom": nom.strip(),
        "entreprise": entreprise.strip(),
        "date": the_date.isoformat(),
        "responsable": responsable,
        "statut": statut,
        "proprietaires": proprietaires,
        "notes": notes.strip(),
    }


@st.cache_data(ttl=120, show_spinner=False)
def leonar_mandate_projects() -> list[dict]:
    return leonar.list_mandate_projects()


def leonar_owner_names(project: dict) -> list[str]:
    return [o.get("first_name") for o in project.get("owners") or [] if o.get("first_name") in PROPRIETAIRES]


def _apply_leonar_import():
    """Remplit le formulaire « Nouveau mandat » avec le projet Leonar choisi."""
    idx = st.session_state.get("new_leonar_choice")
    projects = st.session_state.get("_leonar_import_options") or []
    if idx is None or idx >= len(projects):
        st.session_state.pop("new_leonar_project_id", None)
        return
    p = projects[idx]
    owners = leonar_owner_names(p)
    st.session_state["new_leonar_project_id"] = p["id"]
    st.session_state["new_nom"] = p.get("name", "")
    st.session_state["new_ent"] = leonar.project_client(p)
    try:
        st.session_state["new_date"] = date.fromisoformat((p.get("created_at") or "")[:10])
    except ValueError:
        pass
    st.session_state["new_resp"] = next((o for o in owners if o in RESPONSABLES), RESPONSABLES[0])
    st.session_state["new_proprio"] = owners or [st.session_state["new_resp"]]


def render_create_form(mandats: list[dict]):
    with st.container(border=True):
        st.markdown("#### Nouveau mandat")

        # Import d'un projet Leonar existant (hors projets déjà liés à un mandat)
        linked = {x.get("leonar_project_id") for x in mandats if x.get("leonar_project_id")}
        try:
            options = [p for p in leonar_mandate_projects() if p["id"] not in linked]
        except leonar.LeonarError as e:
            options = []
            st.caption(f"Projets Leonar indisponibles : {e}")
        st.session_state["_leonar_import_options"] = options
        if options:
            st.selectbox(
                "Importer un projet Leonar (facultatif)",
                list(range(len(options))),
                format_func=lambda i: (
                    f"{options[i].get('name', '')} — {leonar.project_client(options[i])} · "
                    f"{fmt_date((options[i].get('created_at') or '')[:10])}"
                ),
                index=None, placeholder="Choisir un projet Leonar pour remplir le formulaire…",
                key="new_leonar_choice", on_change=_apply_leonar_import,
            )
        imported_id = st.session_state.get("new_leonar_project_id")

        with st.form("mandat_create", clear_on_submit=False, border=False):
            meta = mandat_fields("new")

            st.markdown("##### Fiche de poste")
            fc1, fc2 = st.columns(2)
            fiche_file = fc1.file_uploader("Fichier (PDF ou Word)", type=FICHE_TYPES, key="new_fiche_file")
            fiche_text = fc2.text_area(
                "…ou texte de l'annonce copié-collé", height=124, key="new_fiche_text",
                placeholder="Collez ici le texte de la fiche de poste (HelloWork, LinkedIn, Apec…)",
            )

            st.markdown("##### Échanges audio")
            audios = st.file_uploader(
                "Enregistrements (manager, RH…) — chaque audio est transcrit intégralement",
                type=AUDIO_TYPES, accept_multiple_files=True, key="new_audios",
            )

            if imported_id:
                st.caption(":material/link: Le mandat sera lié au projet Leonar importé.")
                create_in_leonar = False
            else:
                create_in_leonar = st.checkbox(
                    "Créer aussi le projet dans Leonar", value=True, key="new_create_leonar",
                    help="Même nom, entreprise en client (fiche créée si besoin) et propriétaires choisis.",
                )
            force = st.checkbox(
                "Créer même si un mandat du même nom existe déjà pour cette entreprise",
                key="new_force",
            )
            with st.container(horizontal=True, gap="small"):
                submitted = st.form_submit_button("Créer le mandat", type="primary")
                cancelled = st.form_submit_button("Annuler")

        if cancelled:
            st.session_state["_mandats_creating"] = False
            st.session_state.pop("new_leonar_project_id", None)
            st.rerun()
        if not submitted:
            return

        if not meta["nom"] or not meta["entreprise"]:
            st.error("Le nom du mandat et l'entreprise sont obligatoires.")
            return
        if not meta["proprietaires"]:
            meta["proprietaires"] = [meta["responsable"]]
        dup = next(
            (
                x for x in mandats
                if x.get("nom", "").strip().lower() == meta["nom"].lower()
                and x.get("entreprise", "").strip().lower() == meta["entreprise"].lower()
            ),
            None,
        )
        if dup and not force:
            st.warning(
                f"Un mandat « {dup['nom']} » chez {dup['entreprise']} existe déjà "
                f"(du {fmt_date(dup.get('date', ''))}). Ouvrez-le pour le compléter, ou cochez "
                "« Créer même si… » pour créer un mandat distinct."
            )
            return

        leonar_warning = None
        if imported_id:
            meta["leonar_project_id"] = imported_id
        elif create_in_leonar:
            try:
                with st.spinner("Création du projet dans Leonar…"):
                    project = leonar.create_project(meta["nom"], meta["entreprise"], meta["proprietaires"])
                meta["leonar_project_id"] = project["id"]
                leonar_mandate_projects.clear()
            except leonar.LeonarError as e:
                leonar_warning = f"Projet Leonar non créé : {e}"

        try:
            with st.spinner("Création du mandat sur le Drive…"):
                mandat_id = store().create_mandat(meta)
                save_fiche(mandat_id, [], fiche_file, fiche_text)
            mandat = {**meta, "id": mandat_id}
            ok = True
            for a in audios or []:
                ok = add_audio(mandat, a) and ok
        except StorageError as e:
            st.error(str(e))
            return

        st.session_state["_mandats_creating"] = False
        st.session_state.pop("new_leonar_project_id", None)
        flash = "Mandat créé." if ok else "Mandat créé — une transcription a échoué, relancez-la depuis la fiche."
        if leonar_warning:
            flash += " " + leonar_warning
        refresh_and_rerun(flash, mandat_id)


# ─────────────────────────────────────────────────────────────────────────────
# Vue liste
# ─────────────────────────────────────────────────────────────────────────────


def view_list(snap: dict):
    mandats = snap["mandats"]
    files = snap["files"]

    with st.container(horizontal=True, vertical_alignment="bottom"):
        with st.container():
            st.title("🗂️ Mandats")
            st.caption("Fiche de poste, échanges audio et transcriptions de chaque mandat — accessibles à toute l'équipe.")
        if st.button("Nouveau mandat", type="primary", icon=":material/add:"):
            st.session_state["_mandats_creating"] = not st.session_state.get("_mandats_creating", False)
            st.rerun()

    if st.session_state.get("_mandats_creating"):
        render_create_form(mandats)

    with st.container(horizontal=True, vertical_alignment="center", gap="small"):
        query = st.text_input(
            "Rechercher", placeholder="Nom du mandat, entreprise…",
            label_visibility="collapsed", icon=":material/search:", key="mandats_q", width=340,
        )
        statut = st.segmented_control(
            "Statut", ["Tous"] + STATUTS, default="Tous",
            label_visibility="collapsed", key="mandats_statut",
        )
        resp = st.selectbox(
            "Responsable", ["Tous les responsables"] + RESPONSABLES,
            label_visibility="collapsed", key="mandats_resp", width=210,
        )
        if st.button("Actualiser", icon=":material/refresh:", type="tertiary"):
            refresh_and_rerun()

    q = (query or "").strip().lower()
    shown = [
        m for m in mandats
        if (not q or q in f"{m.get('nom', '')} {m.get('entreprise', '')} {m.get('responsable', '')}".lower())
        and (statut in (None, "Tous") or m.get("statut") == statut)
        and (resp == "Tous les responsables" or m.get("responsable") == resp)
    ]
    shown.sort(key=sort_key, reverse=True)

    if not mandats:
        st.markdown(
            '<div class="mandat-empty"><b>Aucun mandat pour l\'instant</b>'
            "Cliquez sur « Nouveau mandat » pour enregistrer le premier.</div>",
            unsafe_allow_html=True,
        )
        return

    st.caption(
        f"{len(shown)} mandat{'s' if len(shown) > 1 else ''}"
        + (f" sur {len(mandats)}" if len(shown) != len(mandats) else "")
        + " · du plus récent au plus ancien"
    )

    if not shown:
        st.markdown(
            '<div class="mandat-empty"><b>Aucun résultat</b>Modifiez la recherche ou les filtres.</div>',
            unsafe_allow_html=True,
        )
        return

    for m in shown:
        with st.container(border=True, horizontal=True, vertical_alignment="center", gap="medium"):
            with st.container(gap="small"):
                st.markdown(f"**{md(m.get('nom', ''))}** &nbsp;·&nbsp; :material/apartment: {md(m.get('entreprise', ''))}")
                st.markdown(badges(m))
                st.caption(files_summary(files.get(m["id"], [])))
            if st.button("Ouvrir", key=f"open_{m['id']}", icon=":material/arrow_forward:", icon_position="right"):
                st.query_params["mandat"] = m["id"]
                st.rerun()


# ─────────────────────────────────────────────────────────────────────────────
# Vue détail
# ─────────────────────────────────────────────────────────────────────────────


def view_detail(m: dict, files: list[dict]):
    if st.button("Tous les mandats", type="tertiary", icon=":material/arrow_back:"):
        del st.query_params["mandat"]
        st.session_state.pop("_mandats_editing", None)
        st.rerun()

    st.title(m.get("nom", ""))
    st.markdown(f"#### :material/apartment: {md(m.get('entreprise', ''))}")
    st.markdown(badges(m))

    with st.container(horizontal=True, gap="small"):
        if st.button("Modifier", icon=":material/edit:"):
            st.session_state["_mandats_editing"] = not st.session_state.get("_mandats_editing", False)
            st.rerun()
        link = store().folder_link(m["id"])
        if link:
            st.link_button("Ouvrir dans Drive", link, icon=":material/folder_open:")
        if m.get("leonar_project_id"):
            st.link_button("Ouvrir dans Leonar", leonar.project_url(m["leonar_project_id"]), icon=":material/open_in_new:")
        else:
            render_leonar_link_popover(m, files)
        with st.popover("Supprimer", icon=":material/delete:"):
            st.markdown("Supprimer ce mandat et **tous ses fichiers** ?")
            st.caption("Ils iront dans la corbeille du Drive partagé (récupérables 30 jours).")
            if st.button("Oui, supprimer", type="primary", key="del_mandat"):
                try:
                    store().trash(m["id"])
                except StorageError as e:
                    st.error(str(e))
                else:
                    del st.query_params["mandat"]
                    refresh_and_rerun("Mandat supprimé.")

    if st.session_state.get("_mandats_editing"):
        with st.container(border=True):
            with st.form("mandat_edit", border=False):
                meta = mandat_fields("edit", m)
                with st.container(horizontal=True, gap="small"):
                    saved = st.form_submit_button("Enregistrer", type="primary")
                    cancel = st.form_submit_button("Annuler")
            if cancel:
                st.session_state["_mandats_editing"] = False
                st.rerun()
            if saved:
                if not meta["nom"] or not meta["entreprise"]:
                    st.error("Le nom du mandat et l'entreprise sont obligatoires.")
                else:
                    try:
                        store().save_meta({**m, **meta})
                    except StorageError as e:
                        st.error(str(e))
                    else:
                        st.session_state["_mandats_editing"] = False
                        refresh_and_rerun("Mandat mis à jour.")

    if m.get("proprietaires"):
        st.caption(f"Propriétaires du projet Leonar : {', '.join(m['proprietaires'])}")
    if m.get("notes"):
        with st.container(border=True):
            st.markdown("**:material/sticky_note_2: Notes**")
            st.markdown(md(m["notes"]).replace("\n", "  \n"))
    st.caption(
        f"Créé le {fmt_datetime(m.get('created_at', ''))}"
        + (f" · modifié le {fmt_datetime(m['updated_at'])}" if m.get("updated_at") else "")
    )

    st.divider()
    render_fiche_section(m, files)
    render_audio_section(m, files)
    render_scorecard_section(m, files)


def push_scorecards_to_leonar(project_id: str, scorecards: list[dict]) -> tuple[int, list[str]]:
    """Envoie les scorecards (HTML du Drive → PDF) sur le projet Leonar, en remplaçant les versions précédentes."""
    from utils.pdf_export import html_to_pdf, PdfExportError

    sent, errors = 0, []
    for f in scorecards:
        try:
            pdf = html_to_pdf(fetch_bytes(f["id"]).decode("utf-8", errors="replace"))
            leonar.replace_project_file(project_id, f["name"].replace(".html", ".pdf"), pdf)
            sent += 1
        except (PdfExportError, leonar.LeonarError, StorageError) as e:
            errors.append(f"{f['name']} : {e}")
    return sent, errors


def link_flash(base: str, project_id: str, files: list[dict]) -> str:
    """Après liaison : envoie les scorecards existantes et complète le message affiché."""
    scorecards = [f for f in files if f["kind"] == "scorecard"]
    if not scorecards:
        return base
    with st.spinner("Envoi des scorecards existantes sur Leonar…"):
        sent, errors = push_scorecards_to_leonar(project_id, scorecards)
    msg = base + (f" {sent} scorecard(s) envoyée(s) sur Leonar." if sent else "")
    if errors:
        msg += " Envoi Leonar impossible pour : " + " ; ".join(errors)
    return msg


def render_leonar_link_popover(m: dict, files: list[dict]):
    """Mandat sans projet Leonar : le créer, ou le lier à un projet existant."""
    with st.popover("Lier à Leonar", icon=":material/link:"):
        owners = m.get("proprietaires") or [m.get("responsable") or RESPONSABLES[0]]
        st.markdown("**Créer le projet dans Leonar**")
        st.caption(f"« {m.get('nom', '')} », client {m.get('entreprise', '')}, propriétaires : {', '.join(owners)}.")
        if st.button("Créer dans Leonar", type="primary", key="leonar_create_existing"):
            try:
                with st.spinner("Création du projet dans Leonar…"):
                    project = leonar.create_project(m.get("nom", ""), m.get("entreprise", ""), owners)
                store().save_meta({**m, "leonar_project_id": project["id"]})
            except (leonar.LeonarError, StorageError) as e:
                st.error(str(e))
            else:
                leonar_mandate_projects.clear()
                refresh_and_rerun(link_flash("Projet créé dans Leonar et lié au mandat.", project["id"], files))
        st.divider()
        st.markdown("**…ou lier un projet Leonar existant**")
        try:
            snap = load_snapshot()
            linked = {x.get("leonar_project_id") for x in snap["mandats"] if x.get("leonar_project_id")}
            options = [p for p in leonar_mandate_projects() if p["id"] not in linked]
        except leonar.LeonarError as e:
            st.caption(f"Projets Leonar indisponibles : {e}")
            return
        idx = st.selectbox(
            "Projet Leonar", list(range(len(options))),
            format_func=lambda i: f"{options[i].get('name', '')} — {leonar.project_client(options[i])}",
            index=None, placeholder="Choisir un projet…", key="leonar_link_choice",
        )
        if st.button("Lier ce projet", disabled=idx is None, key="leonar_link_btn"):
            p = options[idx]
            try:
                store().save_meta({**m, "leonar_project_id": p["id"],
                                   "proprietaires": leonar_owner_names(p) or owners})
            except StorageError as e:
                st.error(str(e))
            else:
                refresh_and_rerun(link_flash("Mandat lié au projet Leonar.", p["id"], files))


def render_fiche_section(m: dict, files: list[dict]):
    st.subheader("Fiche de poste")
    fiches = [f for f in files if f["kind"] == "fiche"]
    if not fiches:
        st.caption("Aucune fiche de poste pour l'instant.")
    for f in sorted(fiches, key=lambda f: f["props"].get("format", "")):
        is_text = f["props"].get("format") == "text"
        with st.container(border=True):
            with st.container(horizontal=True, vertical_alignment="center", gap="small"):
                st.markdown(
                    f":material/{'notes' if is_text else 'description'}: **{md(strip_prefix(f['name']))}**"
                    f" &nbsp; :gray[{fmt_size(f['size'])}]",
                    width="stretch",
                )
                st.download_button(
                    "Télécharger", data=lambda fid=f["id"]: store().download(fid),
                    file_name=strip_prefix(f["name"]) if not is_text else "Fiche de poste.txt",
                    mime=f["mime"], on_click="ignore", icon=":material/download:", key=f"dl_{f['id']}",
                )
                with st.popover("", icon=":material/delete:", help="Supprimer cette fiche"):
                    st.markdown("Supprimer cette fiche de poste ?")
                    if st.button("Oui, supprimer", type="primary", key=f"del_{f['id']}"):
                        delete_file(f["id"], "Fiche de poste supprimée.")
            if is_text:
                with st.expander("Lire la fiche"):
                    try:
                        st.text(fetch_bytes(f["id"]).decode("utf-8", errors="replace"))
                    except StorageError as e:
                        st.error(str(e))

    with st.expander("Ajouter ou remplacer la fiche de poste", icon=":material/upload_file:", expanded=not fiches):
        with st.form(f"fiche_{m['id']}", border=False, clear_on_submit=True):
            up = st.file_uploader("Fichier (PDF ou Word)", type=FICHE_TYPES)
            txt = st.text_area("…ou texte de l'annonce copié-collé", height=140)
            st.caption("Un nouveau fichier remplace l'ancien fichier ; un nouveau texte remplace l'ancien texte.")
            if st.form_submit_button("Enregistrer la fiche", type="primary"):
                if up is None and not txt.strip():
                    st.error("Ajoutez un fichier ou collez un texte.")
                else:
                    try:
                        with st.spinner("Envoi sur le Drive…"):
                            save_fiche(m["id"], files, up, txt)
                    except StorageError as e:
                        st.error(str(e))
                    else:
                        refresh_and_rerun("Fiche de poste enregistrée.")


def render_audio_section(m: dict, files: list[dict]):
    st.subheader("Échanges audio")
    audios = sorted([f for f in files if f["kind"] == "audio"], key=lambda f: f["created"])
    transcripts = {f["props"].get("audio_id"): f for f in files if f["kind"] == "transcript"}
    orphans = [t for aid, t in transcripts.items() if aid not in {a["id"] for a in audios}]

    if not audios and not orphans:
        st.caption("Aucun échange audio pour l'instant.")

    for a in audios:
        tr = transcripts.get(a["id"])
        with st.container(border=True):
            status_badge = (
                ":green-badge[:material/check: Transcrit]" if tr
                else ":orange-badge[:material/schedule: Pas encore transcrit]"
            )
            st.markdown(f":material/graphic_eq: **{md(audio_display_name(a))}** &nbsp; {status_badge}")
            st.caption(
                f"{strip_prefix(a['name'])} · {fmt_size(a['size'])} · ajouté le {fmt_datetime(a['created'])}"
            )

            play_key = f"play_{a['id']}"
            with st.container(horizontal=True, gap="small"):
                if st.button("Écouter", icon=":material/play_arrow:", key=f"btn_{play_key}"):
                    st.session_state[play_key] = not st.session_state.get(play_key, False)
                st.download_button(
                    "Audio", data=lambda fid=a["id"]: store().download(fid),
                    file_name=strip_prefix(a["name"]), mime=a["mime"], on_click="ignore",
                    icon=":material/download:", key=f"dl_{a['id']}", help="Télécharger l'audio",
                )
                if tr:
                    st.download_button(
                        "Transcription .txt", data=lambda fid=tr["id"]: store().download(fid),
                        file_name=strip_prefix(tr["name"]), mime="text/plain", on_click="ignore",
                        icon=":material/download:", key=f"dl_{tr['id']}",
                    )
                    st.download_button(
                        "Transcription Word",
                        data=lambda fid=tr["id"], t=audio_display_name(a): transcript_to_docx(
                            store().download(fid).decode("utf-8", errors="replace"),
                            f"Transcription — {t}",
                        ),
                        file_name=os.path.splitext(strip_prefix(tr["name"]))[0] + ".docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        on_click="ignore", icon=":material/download:", key=f"docx_{tr['id']}",
                    )
                else:
                    if st.button("Transcrire", type="primary", icon=":material/transcribe:", key=f"tr_{a['id']}"):
                        st.session_state[f"do_tr_{a['id']}"] = True
                with st.popover("", icon=":material/delete:", help="Supprimer cet audio"):
                    st.markdown("Supprimer cet audio **et sa transcription** ?")
                    if st.button("Oui, supprimer", type="primary", key=f"del_{a['id']}"):
                        try:
                            if tr:
                                store().trash(tr["id"])
                            store().trash(a["id"])
                        except StorageError as e:
                            st.error(str(e))
                        else:
                            refresh_and_rerun("Audio supprimé.")

            if st.session_state.pop(f"do_tr_{a['id']}", False):
                try:
                    with st.spinner("Récupération de l'audio…"):
                        data = store().download(a["id"])
                except StorageError as e:
                    st.error(str(e))
                else:
                    if transcribe_and_store(m, a, data):
                        refresh_and_rerun("Transcription enregistrée.")

            if st.session_state.get(play_key):
                try:
                    with st.spinner("Chargement de l'audio…"):
                        st.audio(fetch_bytes(a["id"]), format=a["mime"] or "audio/mpeg")
                except StorageError as e:
                    st.error(str(e))

            if tr:
                with st.expander("Lire la transcription", icon=":material/article:"):
                    try:
                        text = fetch_bytes(tr["id"]).decode("utf-8", errors="replace")
                    except StorageError as e:
                        st.error(str(e))
                    else:
                        long_text = text.count("\n") > 20
                        with st.container(height=460 if long_text else "content", border=False):
                            render_transcript(text)

    for t in orphans:
        with st.container(border=True, horizontal=True, vertical_alignment="center"):
            st.markdown(
                f":material/article: **{md(strip_prefix(t['name']))}** &nbsp; :gray[audio supprimé]",
                width="stretch",
            )
            st.download_button(
                "Transcription .txt", data=lambda fid=t["id"]: store().download(fid),
                file_name=strip_prefix(t["name"]), mime="text/plain", on_click="ignore",
                icon=":material/download:", key=f"dl_{t['id']}",
            )

    with st.expander("Ajouter un échange audio", icon=":material/mic:", expanded=not audios):
        with st.form(f"audio_{m['id']}", border=False, clear_on_submit=True):
            up = st.file_uploader("Enregistrement audio", type=AUDIO_TYPES)
            label = st.text_input("Intitulé (facultatif)", placeholder="Ex : Échange avec le manager, Échange RH…")
            st.caption("L'audio est enregistré sur le Drive puis transcrit intégralement (locuteurs + horodatage).")
            if st.form_submit_button("Ajouter et transcrire", type="primary", icon=":material/upload:"):
                if up is None:
                    st.error("Choisissez un fichier audio.")
                else:
                    try:
                        ok = add_audio(m, up, label)
                    except StorageError as e:
                        st.error(str(e))
                    else:
                        refresh_and_rerun("Audio ajouté et transcrit." if ok else "Audio ajouté — transcription à relancer.")


def fiche_text_for_scorecard(f: dict) -> str:
    """Texte de la fiche de poste (texte collé, PDF ou Word)."""
    data = fetch_bytes(f["id"])
    if f["props"].get("format") == "text":
        return data.decode("utf-8", errors="replace")
    name = f["name"].lower()
    if name.endswith(".pdf"):
        import pdfplumber
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            return "\n".join(page.extract_text() or "" for page in pdf.pages)
    if name.endswith(".docx"):
        from docx import Document
        return "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
    return ""


def render_scorecard_section(m: dict, files: list[dict]):
    st.subheader("Scorecard")
    scorecards = sorted([f for f in files if f["kind"] == "scorecard"], key=lambda f: f["created"], reverse=True)
    for sc in scorecards:
        with st.container(border=True, horizontal=True, vertical_alignment="center", gap="small"):
            st.markdown(
                f":material/assignment: **{md(strip_prefix(sc['name']))}** &nbsp; "
                f":gray[enregistrée le {fmt_datetime(sc['created'])}]",
                width="stretch",
            )
            st.download_button(
                "Télécharger", data=lambda fid=sc["id"]: store().download(fid),
                file_name=sc["name"], mime="text/html", on_click="ignore",
                icon=":material/download:", key=f"dl_{sc['id']}",
                help="Fichier HTML : ouvrez-le dans le navigateur pour l'imprimer en PDF, "
                     "ou chargez-le dans Dossier Candidature.",
            )
            with st.popover("", icon=":material/delete:", help="Supprimer cette scorecard"):
                st.markdown("Supprimer cette scorecard ?")
                if st.button("Oui, supprimer", type="primary", key=f"del_{sc['id']}"):
                    delete_file(sc["id"], "Scorecard supprimée.")

    transcripts = sorted([f for f in files if f["kind"] == "transcript"], key=lambda f: f["created"])
    fiches = [f for f in files if f["kind"] == "fiche" and not f["name"].lower().endswith(".doc")]
    sources = {f"Transcription — {f['props'].get('label') or strip_prefix(f['name'])}": f for f in transcripts}
    sources.update({f"Fiche de poste — {strip_prefix(f['name'])}": f for f in fiches})

    with st.container(border=True):
        if not sources:
            st.caption("Ajoutez une fiche de poste ou un échange audio transcrit pour générer la scorecard.")
            return
        st.markdown("**Générer la scorecard de ce mandat**")
        st.caption(
            "Ouvre la rubrique Scorecard avec l'entreprise, le responsable et les éléments "
            "ci-dessous déjà chargés. Une fois la version finale validée, elle est enregistrée "
            "dans ce mandat et envoyée sur Leonar."
        )
        chosen = st.multiselect(
            "Éléments à utiliser", list(sources), default=list(sources), key=f"sc_src_{m['id']}",
        )
        if st.button("Ouvrir dans Scorecard", type="primary", icon=":material/arrow_outward:",
                     disabled=not chosen, key=f"sc_go_{m['id']}"):
            parts = []
            try:
                with st.spinner("Préparation des éléments…"):
                    for label in chosen:
                        f = sources[label]
                        text = (
                            fetch_bytes(f["id"]).decode("utf-8", errors="replace")
                            if f["kind"] == "transcript" else fiche_text_for_scorecard(f)
                        )
                        if text.strip():
                            parts.append(f"══════ {label.upper()} ══════\n\n{text.strip()}")
            except Exception as e:
                st.error(f"Impossible de lire les éléments : {e}")
                return
            if not parts:
                st.error("Les éléments choisis ne contiennent pas de texte exploitable.")
                return
            st.session_state["scorecard_mandat"] = {
                "id": m["id"], "nom": m.get("nom", ""), "leonar_project_id": m.get("leonar_project_id"),
            }
            st.session_state["scorecard_prefill"] = {
                "mandat_id": m["id"],
                "mandat_nom": m.get("nom", ""),
                "client": m.get("entreprise", ""),
                "commercial": m.get("responsable", ""),
                "text": "\n\n".join(parts),
                "sources": chosen,
            }
            for k in ["scorecard_html", "scorecard_raw_html", "scorecard_transcription",
                      "scorecard_client", "scorecard_commercial", "scorecard_saved_id", "scorecard_client_input"]:
                st.session_state.pop(k, None)
            st.switch_page("pages/2_Scorecard.py")


def delete_file(file_id: str, flash: str):
    try:
        store().trash(file_id)
    except StorageError as e:
        st.error(str(e))
    else:
        refresh_and_rerun(flash)


# ─────────────────────────────────────────────────────────────────────────────
# Page
# ─────────────────────────────────────────────────────────────────────────────

flash = st.session_state.pop("_mandats_flash", None)
if flash:
    st.toast(flash, icon="✅")

try:
    with st.spinner("Chargement des mandats…"):
        snapshot = load_snapshot()
except StorageError as e:
    st.title("🗂️ Mandats")
    st.error(f"Impossible d'accéder au stockage des mandats : {e}")
    st.stop()
except Exception as e:
    st.title("🗂️ Mandats")
    st.error(f"Erreur inattendue lors du chargement des mandats : {e}")
    st.stop()

selected_id = st.query_params.get("mandat")
selected = next((x for x in snapshot["mandats"] if x["id"] == selected_id), None) if selected_id else None

if selected_id and not selected:
    st.warning("Ce mandat est introuvable (supprimé ?). Retour à la liste.")
    del st.query_params["mandat"]

if selected:
    view_detail(selected, snapshot["files"].get(selected["id"], []))
else:
    view_list(snapshot)
